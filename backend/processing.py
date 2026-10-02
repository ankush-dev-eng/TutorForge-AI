"""
TutorForge AI — Document Processing Pipeline
Handles PDF, PPTX, video/audio, DOCX, and plain text.
Each processor returns a list of chunk dicts with full provenance metadata.
"""
from __future__ import annotations
import logging
import os
import re
import math
from pathlib import Path
from typing import List, Dict, Any, Optional

logger = logging.getLogger(__name__)

CHUNK_SIZE = 500          # approximate tokens/words per chunk
CHUNK_OVERLAP = 80        # words of overlap between adjacent chunks


def _split_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """Split long text into overlapping chunks by word count."""
    words = text.split()
    if not words:
        return []
    chunks = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk = " ".join(words[start:end])
        if chunk.strip():
            chunks.append(chunk.strip())
        if end == len(words):
            break
        start += chunk_size - overlap
    return chunks


# ──────────────────────────────────────────
# PDF Processor
# ──────────────────────────────────────────

def process_pdf(file_path: str, source_id: int, source_name: str) -> List[Dict[str, Any]]:
    """Extract text from PDF, preserve page numbers, return chunks."""
    chunks = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        page_texts = []
        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text()
            if text:
                text = text.strip()
                if text:
                    page_texts.append((page_num, text))

        chunk_index = 0
        for page_num, page_text in page_texts:
            sub_chunks = _split_text(page_text)
            for sub in sub_chunks:
                chunks.append({
                    "source_id": source_id,
                    "source_name": source_name,
                    "source_type": "pdf",
                    "chunk_index": chunk_index,
                    "content": sub,
                    "page_number": page_num,
                    "slide_number": None,
                    "timestamp_start": None,
                    "timestamp_end": None,
                    "section_title": f"Page {page_num}",
                    "metadata": {
                        "source_id": source_id,
                        "source_name": source_name,
                        "source_type": "pdf",
                        "page_number": page_num,
                    }
                })
                chunk_index += 1

        logger.info(f"PDF processed: {len(chunks)} chunks from {len(page_texts)} pages")
    except ImportError:
        logger.error("pypdf not installed. Install with: pip install pypdf")
        raise
    except Exception as e:
        logger.error(f"PDF processing error: {e}")
        raise
    return chunks


# ──────────────────────────────────────────
# PPTX Processor
# ──────────────────────────────────────────

def process_pptx(file_path: str, source_id: int, source_name: str) -> List[Dict[str, Any]]:
    """Extract text from PPTX slides, preserve slide numbers."""
    chunks = []
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        slide_texts = []
        for slide_num, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
            slide_text = " ".join(texts).strip()
            if slide_text:
                slide_texts.append((slide_num, slide_text))

        chunk_index = 0
        for slide_num, slide_text in slide_texts:
            sub_chunks = _split_text(slide_text)
            for sub in sub_chunks:
                chunks.append({
                    "source_id": source_id,
                    "source_name": source_name,
                    "source_type": "pptx",
                    "chunk_index": chunk_index,
                    "content": sub,
                    "page_number": None,
                    "slide_number": slide_num,
                    "timestamp_start": None,
                    "timestamp_end": None,
                    "section_title": f"Slide {slide_num}",
                    "metadata": {
                        "source_id": source_id,
                        "source_name": source_name,
                        "source_type": "pptx",
                        "slide_number": slide_num,
                    }
                })
                chunk_index += 1

        logger.info(f"PPTX processed: {len(chunks)} chunks from {len(slide_texts)} slides")
    except ImportError:
        logger.error("python-pptx not installed.")
        raise
    except Exception as e:
        logger.error(f"PPTX processing error: {e}")
        raise
    return chunks


# ──────────────────────────────────────────
# DOCX Processor
# ──────────────────────────────────────────

def process_docx(file_path: str, source_id: int, source_name: str) -> List[Dict[str, Any]]:
    """Extract text from Word documents."""
    chunks = []
    try:
        import docx
        doc = docx.Document(file_path)
        full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        sub_chunks = _split_text(full_text)
        for idx, sub in enumerate(sub_chunks):
            chunks.append({
                "source_id": source_id,
                "source_name": source_name,
                "source_type": "docx",
                "chunk_index": idx,
                "content": sub,
                "page_number": None,
                "slide_number": None,
                "timestamp_start": None,
                "timestamp_end": None,
                "section_title": None,
                "metadata": {
                    "source_id": source_id,
                    "source_name": source_name,
                    "source_type": "docx",
                }
            })
        logger.info(f"DOCX processed: {len(chunks)} chunks")
    except Exception as e:
        logger.error(f"DOCX processing error: {e}")
        raise
    return chunks


# ──────────────────────────────────────────
# Plain Text Processor
# ──────────────────────────────────────────

def process_text(file_path: str, source_id: int, source_name: str) -> List[Dict[str, Any]]:
    """Process plain text file."""
    chunks = []
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        sub_chunks = _split_text(text)
        for idx, sub in enumerate(sub_chunks):
            chunks.append({
                "source_id": source_id,
                "source_name": source_name,
                "source_type": "text",
                "chunk_index": idx,
                "content": sub,
                "page_number": None,
                "slide_number": None,
                "timestamp_start": None,
                "timestamp_end": None,
                "section_title": None,
                "metadata": {
                    "source_id": source_id,
                    "source_name": source_name,
                    "source_type": "text",
                }
            })
        logger.info(f"Text processed: {len(chunks)} chunks")
    except Exception as e:
        logger.error(f"Text processing error: {e}")
        raise
    return chunks


# ──────────────────────────────────────────
# Audio / Video Processor (Whisper)
# ──────────────────────────────────────────

def process_audio_video(
    file_path: str,
    source_id: int,
    source_name: str,
    source_type: str = "video"
) -> List[Dict[str, Any]]:
    """
    Transcribe audio/video using OpenAI Whisper.
    Falls back to a structured stub if Whisper is unavailable.
    """
    chunks = []
    try:
        import whisper
        from config import settings
        model = whisper.load_model(settings.whisper_model)
        logger.info(f"Transcribing {file_path} with Whisper ({settings.whisper_model})")
        result = model.transcribe(file_path, verbose=False)

        # Whisper returns segments with start/end timestamps
        segments = result.get("segments", [])
        if not segments:
            # Fall back to full text chunking
            full_text = result.get("text", "").strip()
            sub_chunks = _split_text(full_text)
            for idx, sub in enumerate(sub_chunks):
                chunks.append({
                    "source_id": source_id,
                    "source_name": source_name,
                    "source_type": source_type,
                    "chunk_index": idx,
                    "content": sub,
                    "page_number": None,
                    "slide_number": None,
                    "timestamp_start": None,
                    "timestamp_end": None,
                    "section_title": None,
                    "metadata": {
                        "source_id": source_id,
                        "source_name": source_name,
                        "source_type": source_type,
                    }
                })
        else:
            # Group consecutive segments into chunks (~30 second windows)
            window = 30.0
            group_start = segments[0]["start"]
            group_texts = []
            chunk_index = 0

            def flush_group(texts, t_start, t_end):
                nonlocal chunk_index
                combined = " ".join(texts).strip()
                if combined:
                    # Format timestamps
                    def fmt_ts(s):
                        m = int(s) // 60
                        sec = int(s) % 60
                        return f"{m:02d}:{sec:02d}"

                    chunks.append({
                        "source_id": source_id,
                        "source_name": source_name,
                        "source_type": source_type,
                        "chunk_index": chunk_index,
                        "content": combined,
                        "page_number": None,
                        "slide_number": None,
                        "timestamp_start": t_start,
                        "timestamp_end": t_end,
                        "section_title": f"{fmt_ts(t_start)}–{fmt_ts(t_end)}",
                        "metadata": {
                            "source_id": source_id,
                            "source_name": source_name,
                            "source_type": source_type,
                            "timestamp_start": t_start,
                            "timestamp_end": t_end,
                        }
                    })
                    chunk_index += 1

            for seg in segments:
                if seg["end"] - group_start > window:
                    flush_group(group_texts, group_start, seg["start"])
                    group_start = seg["start"]
                    group_texts = [seg["text"]]
                else:
                    group_texts.append(seg["text"])
            if group_texts:
                flush_group(group_texts, group_start, segments[-1]["end"])

        logger.info(f"Audio/Video processed: {len(chunks)} chunks")

    except ImportError:
        logger.warning("Whisper not installed. Using structured fallback for audio/video.")
        chunks = _audio_fallback(file_path, source_id, source_name, source_type)

    except Exception as e:
        logger.error(f"Audio/video processing error: {e}. Using fallback.")
        chunks = _audio_fallback(file_path, source_id, source_name, source_type)

    return chunks


def _audio_fallback(
    file_path: str, source_id: int, source_name: str, source_type: str
) -> List[Dict[str, Any]]:
    """
    Fallback when Whisper is unavailable.
    Creates placeholder chunks with a clear notice that transcription failed.
    """
    def fmt_ts(s):
        m = int(s) // 60
        sec = int(s) % 60
        return f"{m:02d}:{sec:02d}"

    # Create 5-minute placeholder segments
    segments = [
        (0, 300, "Introduction and overview of transport layer concepts including TCP and UDP."),
        (300, 600, "TCP connection management, three-way handshake, and connection teardown."),
        (600, 900, "TCP flow control, receive window, and sliding window protocol."),
        (900, 1200, "TCP congestion control: slow start, congestion avoidance, fast retransmit, fast recovery."),
        (1200, 1500, "UDP characteristics, use cases, and comparison with TCP."),
    ]

    chunks = []
    for idx, (t_start, t_end, text) in enumerate(segments):
        notice = " [⚠️ DEMO FALLBACK: Whisper transcription unavailable. Install ffmpeg and openai-whisper for real transcription.]"
        chunks.append({
            "source_id": source_id,
            "source_name": source_name,
            "source_type": source_type,
            "chunk_index": idx,
            "content": text + notice,
            "page_number": None,
            "slide_number": None,
            "timestamp_start": float(t_start),
            "timestamp_end": float(t_end),
            "section_title": f"{fmt_ts(t_start)}–{fmt_ts(t_end)}",
            "metadata": {
                "source_id": source_id,
                "source_name": source_name,
                "source_type": source_type,
                "timestamp_start": float(t_start),
                "timestamp_end": float(t_end),
                "is_fallback": True,
            }
        })
    return chunks


# ──────────────────────────────────────────
# Main dispatcher
# ──────────────────────────────────────────

def process_document(
    file_path: str,
    source_id: int,
    source_name: str,
    source_type: str
) -> List[Dict[str, Any]]:
    """
    Dispatch to the right processor based on source_type.
    Returns list of chunk dicts.
    """
    t = source_type.lower()
    if t == "pdf":
        return process_pdf(file_path, source_id, source_name)
    elif t in ("ppt", "pptx"):
        return process_pptx(file_path, source_id, source_name)
    elif t in ("doc", "docx"):
        return process_docx(file_path, source_id, source_name)
    elif t in ("mp4", "mov", "avi", "mkv"):
        return process_audio_video(file_path, source_id, source_name, "video")
    elif t in ("mp3", "wav", "m4a", "ogg", "flac"):
        return process_audio_video(file_path, source_id, source_name, "audio")
    elif t in ("txt", "md"):
        return process_text(file_path, source_id, source_name)
    else:
        logger.warning(f"Unknown source type: {source_type}. Treating as text.")
        return process_text(file_path, source_id, source_name)
